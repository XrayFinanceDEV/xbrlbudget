"""Word (.docx) dei report ReportLab: traduce i flowable delle sezioni, non ne riscrive il contenuto.

Spec 2026-09-24 §2. Le sezioni restano l'unica fonte dei testi: un testo corretto nel codice cambia PDF e Word
insieme. Un flowable che il traduttore non conosce fa fallire la traduzione: una sezione che perde un pezzo nel Word
senza dirlo è il difetto da non avere.
"""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.platypus import CondPageBreak, Image, KeepTogether, NextPageTemplate, PageBreak, Paragraph, Spacer, Table

from app.renderers.business_plan import theme
from app.renderers.business_plan.layout import CoverLines  # noqa: F401  (riesportata)

FONT = "Lato"
_ALIGN = {0: WD_ALIGN_PARAGRAPH.LEFT, 1: WD_ALIGN_PARAGRAPH.CENTER, 2: WD_ALIGN_PARAGRAPH.RIGHT,
          4: WD_ALIGN_PARAGRAPH.JUSTIFY, "LEFT": WD_ALIGN_PARAGRAPH.LEFT, "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
          "CENTRE": WD_ALIGN_PARAGRAPH.CENTER, "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT}
_IGNORATI = (NextPageTemplate, CondPageBreak)


def _hex(color) -> str | None:
    if color is None or getattr(color, "alpha", 1) == 0:
        return None
    return color.hexval()[2:].upper().rjust(6, "0")


def _font(run, name: str, size: float, color, bold: bool) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold or "Bold" in (name or "")
    hx = _hex(color)
    if hx:
        run.font.color.rgb = RGBColor.from_string(hx)


def _new_paragraph(container, state):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(0), Pt(0)
    state["last"] = p
    return p


def _paragraph(container, para: Paragraph, state) -> None:
    p = _new_paragraph(container, state)
    st = para.style
    p.alignment = _ALIGN.get(st.alignment, WD_ALIGN_PARAGRAPH.LEFT)
    if st.leftIndent:
        p.paragraph_format.left_indent = Pt(st.leftIndent)
    for fr in para.frags:
        if getattr(fr, "lineBreak", False):
            p.add_run().add_break()
            continue
        if fr.text:
            _font(p.add_run(fr.text), fr.fontName, fr.fontSize, fr.textColor, False)


def _norm(i: int, n: int) -> int:
    return i + n if i < 0 else i


def _cells(cmd, nrows, ncols):
    (c0, r0), (c1, r1) = cmd[1], cmd[2]
    c0, c1, r0, r1 = _norm(c0, ncols), _norm(c1, ncols), _norm(r0, nrows), _norm(r1, nrows)
    return [(r, c) for r in range(r0, min(r1, nrows - 1) + 1) for c in range(c0, min(c1, ncols - 1) + 1)]


# Word è severo sull'ordine dei figli di w:tcPr e w:tcBorders (schema WordprocessingML): fuori ordine il file può
# aprirsi come «contenuto illeggibile». python-docx e LibreOffice lo tollerano, quindi nessun test di apertura lo vede.
_DOPO_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
_ORDINE_LATI = ("top", "left", "bottom", "right")


def _shade(cell, hx: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:shd"))
    if old is not None:
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), hx)
    tcPr.insert_element_before(shd, *_DOPO_SHD)


def _border(cell, side: str, width: float, hx: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.insert_element_before(borders, "w:shd", *_DOPO_SHD)
    old = borders.find(qn(f"w:{side}"))  # due comandi sullo stesso lato: vince l'ultimo, come nel PDF
    if old is not None:
        borders.remove(old)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single"), b.set(qn("w:sz"), str(max(2, round(width * 8)))), b.set(qn("w:color"), hx)
    dopo = [x for x in _ORDINE_LATI[_ORDINE_LATI.index(side) + 1:]]
    succ = next((borders.find(qn(f"w:{x}")) for x in dopo if borders.find(qn(f"w:{x}")) is not None), None)
    if succ is None:
        borders.append(b)
    else:
        succ.addprevious(b)


def _widths(wt, widths: list) -> None:
    """Larghezze su griglia, celle e tabella: Word legge tcW, LibreOffice e Google Docs la griglia."""
    for g, w in zip(wt._tbl.tblGrid.gridCol_lst, widths):
        g.w = Pt(w)
    for row in wt.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Pt(w)
    tblW = wt._tbl.tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        wt._tbl.tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa"), tblW.set(qn("w:w"), str(round(sum(widths) * 20)))


_LATI = {"LINEABOVE": "top", "LINEBELOW": "bottom", "LINEBEFORE": "left", "LINEAFTER": "right"}


def _table(container, t: Table, state) -> None:
    nrows, ncols = t._nrows, t._ncols
    wt = container.add_table(rows=nrows, cols=ncols)
    wt.alignment = WD_TABLE_ALIGNMENT.CENTER
    wt.autofit = False
    widths = [w if isinstance(w, (int, float)) else None for w in t._colWidths]
    free = [i for i, w in enumerate(widths) if w is None]
    if free:
        rest = max(theme.CW - sum(w for w in widths if w), 20)
        for i in free:
            widths[i] = rest / len(free)
    _widths(wt, widths)
    for cmd in t._bkgrndcmds:
        if cmd[0] == "BACKGROUND" and _hex(cmd[3]):
            for r, c in _cells(cmd, nrows, ncols):
                _shade(wt.cell(r, c), _hex(cmd[3]))
    for cmd in t._linecmds:
        side = _LATI.get(cmd[0])
        if side and _hex(cmd[4]):
            for r, c in _cells(cmd, nrows, ncols):
                _border(wt.cell(r, c), side, cmd[3], _hex(cmd[4]))
    for r in range(min(t.repeatRows or 0, nrows)):
        trPr = wt.rows[r]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))
    anchors = {}
    for cmd in t._spanCmds:
        cells = _cells(cmd, nrows, ncols)
        if len(cells) > 1:
            (r0, c0), (r1, c1) = cells[0], cells[-1]
            wt.cell(r0, c0).merge(wt.cell(r1, c1))
            anchors.update({rc: rc == (r0, c0) for rc in cells})
    for r, row in enumerate(t._cellvalues):
        for c, value in enumerate(row):
            if anchors.get((r, c)) is False:
                continue
            cell = wt.cell(r, c)
            _cell(cell, value, t._cellStyles[r][c], state)
    state["last"] = None


def _cell(cell, value, cs, state) -> None:
    first = cell.paragraphs[0]
    inner = {"last": None}
    items = value if isinstance(value, (list, tuple)) else [value]
    for v in items:
        if isinstance(v, str):
            if v:
                p = _new_paragraph(cell, inner)
                p.alignment = _ALIGN.get(cs.alignment, WD_ALIGN_PARAGRAPH.LEFT)
                _font(p.add_run(v), cs.fontname, cs.fontsize, cs.color, False)
        elif v is not None:
            translate(cell, [v], inner)
    if len(cell.paragraphs) > 1 and not first.text and not first.runs:
        first._element.getparent().remove(first._element)  # la cella nasce con un paragrafo vuoto


def _image(container, img: Image, state) -> None:
    png = getattr(img, "png", None)
    if png is None:
        raise TypeError("immagine senza PNG: i grafici passano da layout.chart")
    p = _new_paragraph(container, state)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(io.BytesIO(png), width=Pt(img.drawWidth))


def translate(container, flowables, state: dict | None = None) -> None:
    """Scrive i flowable in `container` (il documento o una cella)."""
    state = state if state is not None else {"last": None}
    for f in flowables:
        if isinstance(f, Paragraph):
            _paragraph(container, f, state)
        elif isinstance(f, Table):
            _table(container, f, state)
        elif isinstance(f, Image):
            _image(container, f, state)
        elif isinstance(f, Spacer):
            if state.get("last") is not None:
                pf = state["last"].paragraph_format
                pf.space_after = Pt((pf.space_after.pt if pf.space_after else 0) + f.height)
        elif isinstance(f, PageBreak):
            _new_paragraph(container, state).add_run().add_break(WD_BREAK.PAGE)
            state["last"] = None
        elif isinstance(f, KeepTogether):
            translate(container, f._content, state)
        elif callable(getattr(f, "make", None)):  # tabella che nel PDF si stringe per stare in pagina
            translate(container, [f.make(1.0)], state)
        elif isinstance(f, _IGNORATI) or type(f).__name__ == "SectionStart":
            continue
        else:
            raise TypeError(f"flowable non tradotto in Word: {type(f).__name__}")


def docx_text(data: bytes) -> list:
    """I testi di tutti i paragrafi del corpo, celle annidate comprese, con gli spazi normalizzati."""
    doc = Document(io.BytesIO(data))
    out = []
    for p in doc.element.body.iter(qn("w:p")):
        txt = "".join(t.text or "" for t in p.iter(qn("w:t")))
        txt = re.sub(r"\s+", " ", txt.replace("\xa0", " ")).strip()
        if txt:
            out.append(txt)
    return out


def _page_field(p) -> None:
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run()
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        r._element.append(el)
        _font(r, "", 7.8, None, False)


def _two_sided(p, left: str, right: str, size: float, color: str, bold_left: bool, page: bool = False) -> None:
    # gli stili Header/Footer del modello hanno stop a 4680 e 9360 twip: la tabulazione finirebbe lì, non al margine
    p.style.paragraph_format.tab_stops.clear_all()
    p.paragraph_format.tab_stops.add_tab_stop(Pt(theme.CW), WD_TAB_ALIGNMENT.RIGHT)
    c = RGBColor.from_string(color[1:])
    r = p.add_run(left); _font(r, "", size, None, bold_left); r.font.color.rgb = c
    r = p.add_run("\t" + right); _font(r, "", size, None, False); r.font.color.rgb = c
    if page:
        _page_field(p)


def _cover(doc, cover: CoverLines) -> None:
    band = doc.add_table(rows=1, cols=1)
    cell = band.cell(0, 0)
    _widths(band, [theme.CW])
    _shade(cell, theme.NAVY[1:])
    _border(cell, "bottom", 5.6, theme.TEAL[1:])
    white = RGBColor(0xFF, 0xFF, 0xFF)
    first = True
    for text, size, bold, color, after in ((cover.eyebrow, 10, True, white, 18), (cover.name, 30, True, white, 6),
                                          (cover.title, 21, False, white, 14)) + tuple(
            (line, 11, False, RGBColor.from_string(theme.COVER_SUB[1:]), 4) for line in cover.lines):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text); _font(r, "", size, None, bold); r.font.color.rgb = color
    doc.add_paragraph().paragraph_format.space_after = Pt(12)


def build(sections: list, *, cover: CoverLines, header_left: str, header_right: str, footer: str,
          title: str) -> bytes:
    """Documento A4 coi margini del PDF: fascia di copertina, poi una sezione per pagina."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(9.3)
    normal.paragraph_format.space_after = Pt(0)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(theme.PAGE_W), Pt(theme.PAGE_H)
    sec.left_margin = sec.right_margin = Pt(theme.LM)
    sec.top_margin, sec.bottom_margin = Pt(55), Pt(40)
    sec.header_distance, sec.footer_distance = Pt(14), Pt(16)
    sec.different_first_page_header_footer = True
    _two_sided(sec.header.paragraphs[0], header_left, header_right, 9, theme.NAVY, True)
    _two_sided(sec.footer.paragraphs[0], footer, "", 7.8, theme.MUTED, False, page=True)
    _two_sided(sec.first_page_footer.paragraphs[0], "Riservato e confidenziale", "", 7.8, theme.MUTED, False)
    _cover(doc, cover)
    for n, flowables in enumerate(sections):
        if n > 0:
            doc.add_page_break()
        translate(doc, flowables)
    doc.core_properties.title, doc.core_properties.author = title, ""
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
