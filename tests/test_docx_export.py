"""Traduttore ReportLab → Word: markup, tabelle, immagini, sconosciuti (spec 2026-09-24 §2.1)."""
import io
import pytest
from docx import Document
from reportlab.platypus import Flowable, Image, PageBreak, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.colors import HexColor

from app.renderers import docx_export as dx
from app.renderers.business_plan import layout, theme

theme.register_fonts()  # le Paragraph si costruiscono prima di _doc


def _doc(flowables):
    theme.register_fonts()
    d = Document()
    dx.translate(d, flowables)
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def test_markup_del_paragrafo():
    p = Paragraph("<font color='#2a9d8f'>•</font>&nbsp;&nbsp;<b>Ricavi.</b> testo &amp; altro<br/>riga due",
                  layout.ST["body"])
    doc = Document(io.BytesIO(_doc([p])))
    par = doc.paragraphs[-1]
    assert par.text == "•\xa0\xa0Ricavi. testo & altro\nriga due"
    bold = [r.text for r in par.runs if r.bold]
    assert bold == ["Ricavi."]
    assert str(par.runs[0].font.color.rgb) == "2A9D8F"


def test_tabella_sfondi_intestazione_e_annidata():
    t = layout.fin_table(["2025", "2026"], [("Ricavi", ["1.000", "2.000"], ""), ("EBITDA", ["10", "20"], "hl")])
    tiles = layout.tiles([("€ 1 mln", "Ricavi", "+5%"), ("€ 2 mln", "EBITDA", "+1%")])
    doc = Document(io.BytesIO(_doc([t, tiles])))
    ft = doc.tables[0]
    assert [c.text for c in ft.rows[0].cells] == ["Voce (euro)", "2025", "2026"]
    assert ft.rows[0]._tr.trPr is not None and ft.rows[0]._tr.trPr.find(dx.qn("w:tblHeader")) is not None
    shd = ft.rows[2].cells[0]._tc.tcPr.find(dx.qn("w:shd"))
    assert shd.get(dx.qn("w:fill")).lower() == theme.HL[1:]
    testi = dx.docx_text(_doc([tiles]))
    assert "€ 1 mln" in testi and "EBITDA" in testi and "+1%" in testi


def test_immagine_e_interruzioni():
    from app.renderers.infrannuale import charts
    img = layout.chart(charts.crisi(["a"], [3], 14, ["D · Crisi"]), 2.4)
    assert img.png.startswith(b"\x89PNG")
    doc = Document(io.BytesIO(_doc([img, Spacer(0, 5), PageBreak(), Paragraph("dopo", layout.ST["body"])])))
    assert len(doc.inline_shapes) == 1
    assert abs(doc.inline_shapes[0].width.pt - img.drawWidth) < 1


def test_flowable_sconosciuto_non_passa_in_silenzio():
    class Strano(Flowable):
        def wrap(self, *a): return 0, 0
    with pytest.raises(TypeError, match="Strano"):
        _doc([Strano()])


_TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
         "tcFitText", "vAlign", "hideMark"]
_LATI = ["top", "start", "left", "bottom", "end", "right", "insideH", "insideV", "tl2br", "tr2bl"]


def _campione():
    t = layout.fin_table(["2025"], [("Ricavi", ["1"], ""), ("EBITDA", ["2"], "hl"), ("Totale", ["3"], "bold")])
    t.setStyle(TableStyle([("LINEBELOW", (0, 1), (-1, 1), 0.5, HexColor("#000000")),
                           ("LINEABOVE", (0, 1), (-1, 1), 0.5, HexColor("#000000"))]))
    return dx.build([[t, layout.panel("Lettura", [("", "testo")])]],
                    cover=layout.CoverLines("E", "Nome", "Titolo", ("riga",)), header_left="A", header_right="B",
                    footer="C", title="T")


def test_ordine_dello_schema_nelle_celle():
    """Rilievo 1 della revisione: Word rifiuta tcPr e tcBorders con i figli fuori dall'ordine dello schema."""
    from docx.oxml.ns import qn
    doc = Document(io.BytesIO(_campione()))
    for tc in doc.element.body.iter(qn("w:tc")):
        pr = tc.find(qn("w:tcPr"))
        if pr is None:
            continue
        figli = [c.tag.split("}")[1] for c in pr]
        assert figli == sorted(figli, key=_TCPR.index), figli
        b = pr.find(qn("w:tcBorders"))
        if b is not None:
            lati = [c.tag.split("}")[1] for c in b]
            assert lati == sorted(lati, key=_LATI.index) and len(lati) == len(set(lati)), lati


def test_intestazione_e_pie_alla_destra_del_margine():
    """Rilievo 2: gli stop degli stili Header/Footer (centro 4680, destra 9360) catturavano la tabulazione."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    doc = Document(io.BytesIO(_campione()))
    for p in (doc.sections[0].header.paragraphs[0], doc.sections[0].footer.paragraphs[0]):
        stops = [(round(s.position.pt), s.alignment) for s in p.paragraph_format.tab_stops
                 if s.alignment != WD_TAB_ALIGNMENT.CLEAR]
        stili = [(round(s.position.pt), s.alignment) for s in p.style.paragraph_format.tab_stops]
        attivi = [s for s in stops + stili if s[0] < round(theme.CW)]
        assert attivi == [], attivi


def test_griglia_uguale_alle_larghezze():
    """Rilievo 3: tblGrid a colonne uguali contro tcW reali; LibreOffice e Google Docs seguono la griglia."""
    from docx.oxml.ns import qn
    doc = Document(io.BytesIO(_campione()))
    for tbl in doc.element.body.iter(qn("w:tbl")):
        grid = [int(g.get(qn("w:w"))) for g in tbl.find(qn("w:tblGrid"))]
        first = tbl.find(qn("w:tr"))
        tcw = [int(tc.find(qn("w:tcPr")).find(qn("w:tcW")).get(qn("w:w"))) for tc in first.findall(qn("w:tc"))]
        assert grid == tcw, (grid, tcw)
