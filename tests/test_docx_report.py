"""Il Word contiene tutto il PDF: ogni riga di testo del PDF, tolte intestazioni, piè e numeri di pagina, si trova
nel .docx; i grafici ci sono tutti (spec 2026-09-24 §5)."""
import io
import re

import fitz
import pytest
from docx import Document

from app.renderers.docx_export import docx_text
from tests.test_inf_data import _data, db  # noqa: F401  (fixture condivisa)


def _norm(s):
    return re.sub(r"\s+", " ", s.replace("\xa0", " ")).strip()


def _manca(pdf: bytes, docx: bytes, salta: set) -> list:
    corpus = "\n".join(docx_text(docx))
    mancanti = []
    for n, page in enumerate(fitz.open(stream=pdf, filetype="pdf"), start=1):
        for line in page.get_text().splitlines():
            t = _norm(line)
            if not t or t.isdigit() or t in salta or t.endswith("…"):
                continue
            if t not in corpus:
                mancanti.append((n, t))
    return mancanti


def _controlla(pdf, docx, page_texts, n_grafici):
    left, right, foot = page_texts
    salta = {_norm(left), _norm(right), _norm(foot), "Riservato e confidenziale"}
    assert _manca(pdf, docx, salta) == []
    doc = Document(io.BytesIO(docx))
    assert len(doc.inline_shapes) == n_grafici
    assert doc.sections[0].different_first_page_header_footer
    assert _norm(right) in _norm(doc.sections[0].header.paragraphs[0].text)


@pytest.mark.parametrize("scenario", (17, 23, 3, 5))
def test_infrannuale_word_completo(db, scenario):  # noqa: F811
    from app.renderers.infrannuale.document import page_texts, render_infrannuale, render_infrannuale_docx
    try:
        d = _data(db, scenario)
    except ValueError as e:
        pytest.skip(f"scenario {scenario} incompleto nel DB locale: {e}")
    pdf = render_infrannuale(d)
    n = sum(len(p.get_images()) for p in fitz.open(stream=pdf, filetype="pdf"))
    _controlla(pdf, render_infrannuale_docx(d), page_texts(d), n)


def test_business_plan_word_completo(db):  # noqa: F811
    from database.models import BudgetScenario
    from app.renderers.business_plan import from_report, render_business_plan
    from app.renderers.business_plan.document import page_texts, render_business_plan_docx
    from app.services.final_report_service import assemble_final_report
    sc = db.query(BudgetScenario).filter(BudgetScenario.company_id == 575,
                                         BudgetScenario.scenario_type != "infrannuale").first()
    if sc is None:
        pytest.skip("nessuno scenario budget per AMBIENTA")
    d = from_report(assemble_final_report(db, 575, sc.id, schema_version=2), draft=True)
    pdf = render_business_plan(d)
    n = sum(len(p.get_images()) for p in fitz.open(stream=pdf, filetype="pdf"))
    docx = render_business_plan_docx(d)
    _controlla(pdf, docx, page_texts(d), n)
    assert "BOZZA" in Document(io.BytesIO(docx)).sections[0].header.paragraphs[0].text


def test_indice_senza_numeri_di_pagina(db):  # noqa: F811
    from app.renderers.infrannuale.document import render_infrannuale_docx
    doc = Document(io.BytesIO(render_infrannuale_docx(_data(db, 17))))
    indice = next(t for t in doc.tables if len(t.columns) == 3 and any("Segnali extracontabili" in c.text for c in t.columns[1].cells))
    assert all(not c.text.strip() for c in indice.columns[2].cells)
